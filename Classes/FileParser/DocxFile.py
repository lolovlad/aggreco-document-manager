from pathlib import Path
from re import findall

from docx import Document
from docx.table import Table

from ..Models.Fileschame import *

from .Cell import Cell
from .Table import Table as CustomTable


from .File import File


class DocxFile(File):
    def __init__(self, path: Path):
        self.__file: Document = Document(path)

    def get_all_parser_table_in_file(self, re: str) -> list[CustomTable]:
        tables = self.__file.tables
        custom_tables = []
        for table in tables:
            if self.is_table_schema(table, re):
                custom_table = self.__create_table(table)
                custom_tables.append(custom_table)
        return custom_tables

    def render(self, schemas: dict):
        pass

    def __create_table(self, table: Table) -> CustomTable:
        custom_table = CustomTable(len(table.rows), len(table.columns))
        cells = self.get_list_cells(table)
        cells = self.__create_list_cells(list(cells))
        for cell in cells:
            custom_table.add_cell(cell)
        return custom_table

    def is_table_schema(self, table: Table, re: str) -> bool:
        for k in table._cells:
            if len(findall(re, k.text)) > 0:
                return True
        return False

    def get_list_cells(self, table: Table):
        cells = set()
        for y, row in enumerate(table.columns):
            for x, cell in enumerate(row.cells):
                cell_loc = self.__get_coord_cell(cell)
                cells.add((cell.text, *cell_loc))
        return cells
    def __get_coord_cell(self, cell):
        tc = cell._tc
        try:
            return (tc.top, tc.bottom, tc.left, tc.right)
        except ValueError:

            return (tc.top, self.__get__bottom_coord_cell(cell._parent, cell, (tc.top, tc.left)), tc.left, tc.right)

    def __get__bottom_coord_cell(self, table: Table, cell, coord) -> int:
        text_cell = cell.text
        bottom_coord = coord[0]
        next_text_cell = table.cell(bottom_coord, coord[1])
        while text_cell == next_text_cell.text:
            bottom_coord += 1
            next_text_cell = table.cell(bottom_coord, coord[1])
        return bottom_coord

    def __create_list_cells(self, cells: list) -> list[Cell]:
        return [Cell(top=cell[1], bottom=cell[2], left=cell[3], right=cell[4], text=cell[0]) for cell in cells]

    def __create_list_cells_schames(self, cells: list[Cell], re: str):
        cells_schemas = [cell.get_schemas() for cell in cells]
        return cells_schemas